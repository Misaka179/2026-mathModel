"""
将Q4-2结果填入官方模板 result4-2.xlsx
"""
import json
import openpyxl
from datetime import datetime, time
from pathlib import Path

# 路径配置
JSON_PATH = r'C:\Users\30292\Documents\shumo\2026-mathModel\questions\Q4\05_results\Q4-2\Q4_2_backtest_results_v2.json'
TEMPLATE_PATH = r'C:\Users\30292\Documents\shumo\2026-mathModel\C题\附件\附件5\result4-2.xlsx'
OUTPUT_PATH = r'C:\Users\30292\Documents\shumo\2026-mathModel\questions\Q4\05_results\Q4-2\result4-2_new.xlsx'

def load_json_results():
    """加载JSON结果"""
    print("加载JSON结果...")
    with open(JSON_PATH, 'r', encoding='utf-8') as f:
        data = json.load(f)
    print(f"[OK] 加载完成: {len(data['results'])}天数据")
    return data

def fill_planned_energy(ws, results):
    """填充计划购电量工作表"""
    print("\n填充[计划购电量]工作表...")

    for day_idx, day_data in enumerate(results):
        row = day_idx + 2  # 从第2行开始（第1行是表头）

        # A列：日期（已有）
        # B到EO列（第2到145列）：144个时段的计划购电量
        for t in range(144):
            col = t + 2  # B列是第2列
            E_plan_value = day_data['E_plan'][t]
            ws.cell(row=row, column=col, value=E_plan_value)

        # EP列（第146列）：全天购电总量
        total_plan = sum(day_data['E_plan'])
        ws.cell(row=row, column=146, value=total_plan)

        # EQ列（第147列）：全天购电成本
        ws.cell(row=row, column=147, value=day_data['plan_cost'])

        if (day_idx + 1) % 50 == 0:
            print(f"  已处理 {day_idx + 1}/{len(results)} 天")

    print(f"[OK] 完成：{len(results)}天 × 144时段")

def fill_storage_energy(ws, results):
    """填充储能电量工作表"""
    print("\n填充[储能电量]工作表...")

    # 每天6行（6个时段：0-4, 4-8, 8-12, 12-16, 16-20, 20-24）
    # 时段划分：每个时段包含24个10分钟步长
    time_slots = [
        (0, 24, '0:00-4:00'),    # t=0~23
        (24, 48, '4:00-8:00'),   # t=24~47
        (48, 72, '8:00-12:00'),  # t=48~71
        (72, 96, '12:00-16:00'), # t=72~95
        (96, 120, '16:00-20:00'), # t=96~119
        (120, 144, '20:00-24:00') # t=120~143
    ]

    row = 2  # 从第2行开始

    for day_idx, day_data in enumerate(results):
        date = datetime.strptime(day_data['date'], '%Y-%m-%d')

        for slot_idx, (t_start, t_end, time_range) in enumerate(time_slots):
            # A列：日期（只在第一个时段显示）
            if slot_idx == 0:
                ws.cell(row=row, column=1, value=date)

            # B列：时间段
            ws.cell(row=row, column=2, value=time_range)

            # C列：充电量（kWh）= sum(charge[t_start:t_end]) * (1/6)
            charge_energy = sum(day_data['charge'][t_start:t_end]) / 6.0
            ws.cell(row=row, column=3, value=charge_energy)

            # D列：放电量（kWh）= sum(discharge[t_start:t_end]) * (1/6)
            discharge_energy = sum(day_data['discharge'][t_start:t_end]) / 6.0
            ws.cell(row=row, column=4, value=discharge_energy)

            # E列：时段开始时间
            if slot_idx == 0:
                ws.cell(row=row, column=5, value=time(0, 0))

            # F列：时段结束SOC（取时段最后一个时刻的SOC）
            end_soc = day_data['soc'][t_end - 1]
            ws.cell(row=row, column=6, value=end_soc)

            row += 1

        if (day_idx + 1) % 50 == 0:
            print(f"  已处理 {day_idx + 1}/{len(results)} 天")

    print(f"[OK] 完成：{len(results)}天 × 6时段 = {len(results) * 6}行")

def fill_emergency_energy(ws, results):
    """填充紧急购电量工作表"""
    print("\n填充[紧急购电量]工作表...")

    # 与计划购电量格式相同
    for day_idx, day_data in enumerate(results):
        row = day_idx + 2

        # B到EO列：144个时段的紧急购电量
        for t in range(144):
            col = t + 2
            emergency_value = day_data['emergency'][t]
            ws.cell(row=row, column=col, value=emergency_value)

        # EP列：全天紧急购电总量
        total_emergency = sum(day_data['emergency'])
        ws.cell(row=row, column=146, value=total_emergency)

        # EQ列：全天紧急购电成本
        ws.cell(row=row, column=147, value=day_data['emergency_cost'])

        if (day_idx + 1) % 50 == 0:
            print(f"  已处理 {day_idx + 1}/{len(results)} 天")

    print(f"[OK] 完成：{len(results)}天 × 144时段")

def main():
    print("=" * 60)
    print("Q4-2 结果填入官方模板")
    print("=" * 60)

    # 1. 加载数据
    data = load_json_results()
    results = data['results']

    # 2. 加载模板
    print(f"\n加载模板: {TEMPLATE_PATH}")
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    print(f"[OK] 模板加载成功，工作表: {wb.sheetnames}")

    # 3. 填充各工作表
    sheet_names = wb.sheetnames

    # 计划购电量
    if len(sheet_names) > 0:
        fill_planned_energy(wb[sheet_names[0]], results)

    # 储能电量
    if len(sheet_names) > 1:
        fill_storage_energy(wb[sheet_names[1]], results)

    # 紧急购电量
    if len(sheet_names) > 2:
        fill_emergency_energy(wb[sheet_names[2]], results)

    # 4. 保存结果
    print(f"\n保存结果: {OUTPUT_PATH}")
    Path(OUTPUT_PATH).parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print("[OK] 保存成功")

    print("\n" + "=" * 60)
    print("完成！")
    print(f"输出文件: {OUTPUT_PATH}")
    print("=" * 60)

if __name__ == "__main__":
    main()
